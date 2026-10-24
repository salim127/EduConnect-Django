from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import urlsplit, unquote
import os

from django.test import SimpleTestCase
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from pypdf import PdfReader

from .document_export import browser_pdf
from .office import export_course


class DocumentExportTests(SimpleTestCase):
    def setUp(self):
        self.course = SimpleNamespace(titre='Cours test', description='Description', chapitres=[], contenu='''
            <h2 style="text-align:center;color:#ff0000;font-size:18pt">Titre enrichi</h2>
            <p>Premier <strong>mot</strong><br>Seconde ligne avec <em>italique</em>.</p>
            <ol><li>Étape un</li><li>Étape deux</li></ol>
            <table><tr><th>Technologie</th><th>Rôle</th></tr><tr><td><b>HTML</b></td><td>Structure</td></tr></table>
            <p><img src="https://example.com/image.png" alt="Illustration"></p>''')
        image = BytesIO()
        Image.new('RGB', (600, 200), 'blue').save(image, 'PNG')
        self.image = image.getvalue()

    def test_word_keeps_runs_paragraphs_tables_and_images(self):
        with patch('pedagogie.export_layout.image_data', return_value=self.image):
            document = Document(export_course(self.course, 'docx'))
        title = next(p for p in document.paragraphs if p.text == 'Titre enrichi')
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(title.runs[0].font.size.pt, 18)
        self.assertEqual(str(title.runs[0].font.color.rgb), 'FF0000')
        paragraph = next(p for p in document.paragraphs if p.text.startswith('Premier'))
        self.assertEqual(paragraph.text, 'Premier mot\nSeconde ligne avec italique.')
        self.assertTrue(next(run for run in paragraph.runs if run.text == 'mot').bold)
        self.assertTrue(next(run for run in paragraph.runs if run.text == 'italique').italic)
        self.assertTrue(any(p.text == '1. Étape un' for p in document.paragraphs))
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(1, 0).text, 'HTML')
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertFalse(any(p.text and not p.text.strip() for p in document.paragraphs))

    def test_pdf_fallback_flows_short_course_without_reserved_slide_spaces(self):
        with patch('pedagogie.document_export.browser_pdf', return_value=False), patch('pedagogie.export_layout.image_data', return_value=self.image):
            document = PdfReader(export_course(self.course, 'pdf'))
        self.assertEqual(len(document.pages), 1)
        text = document.pages[0].extract_text()
        for expected in ('Titre enrichi', 'Premier', 'Seconde ligne', 'HTML', 'Structure'):
            self.assertIn(expected, text)
        self.assertNotIn('_x000D_', text)
        self.assertTrue(document.pages[0].images)

    def test_browser_pdf_preserves_html_formats_and_embeds_images(self):
        observed = {}

        def render(command, **kwargs):
            uri = urlsplit(command[-1])
            path = unquote(uri.path)
            if os.name == 'nt':
                path = path.lstrip('/')
            observed['html'] = Path(path).read_text(encoding='utf-8')
            target = next(value.split('=', 1)[1] for value in command if value.startswith('--print-to-pdf='))
            Path(target).write_bytes(b'%PDF-1.7\nmock')
            return SimpleNamespace(returncode=0)

        with patch.dict(os.environ, {'COURSE_PDF_BROWSER': __file__}), patch('pedagogie.document_export.subprocess.run', side_effect=render), patch('pedagogie.export_layout.image_data', return_value=self.image):
            output = BytesIO()
            self.assertTrue(browser_pdf(self.course, output))
        html = observed['html']
        self.assertIn('text-align:center', html)
        self.assertIn('<strong>mot</strong><br/>', html)
        self.assertIn('data:image/png;base64,', html)
        self.assertNotIn('src="https://example.com', html)
        self.assertNotIn('break-before:page', html)
        self.assertTrue(output.getvalue().startswith(b'%PDF'))
