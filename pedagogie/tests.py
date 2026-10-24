from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from .models import Cours, Exercice
from .models import SupportCours
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from io import BytesIO
import tempfile
from uuid import uuid4
from django.contrib.auth.models import Group
from accounts.roles import ensure_system_roles, LEARNER


class MultiQuestionTests(TestCase):
    def setUp(self):
        ensure_system_roles()
        self.user = get_user_model().objects.create_superuser(username='teacher', password='test')
        self.student = get_user_model().objects.create_user(username='student')
        self.student.groups.add(Group.objects.get(name=LEARNER))
        self.client.force_login(self.user)
        self.course = Cours.objects.create(titre='Cours', enseignant=self.user)
        self.data = dict(titre='Quiz', cours=self.course.pk, type_exercice='qcm', difficulte='facile', question='First?', propositions='A\nB', bonne_reponse='A')
        self.data.update({'questions-TOTAL_FORMS': '1', 'questions-INITIAL_FORMS': '0', 'questions-0-question': 'Second?', 'questions-0-propositions': 'C\nD', 'questions-0-bonne_reponse': 'D'})

    def test_create_edit_and_grade(self):
        response = self.client.post(reverse('pedagogie:exercice_create'), self.data)
        self.assertEqual(response.status_code, 302)
        quiz = Exercice.objects.get()
        self.assertEqual(len(quiz.questions_supplementaires), 1)
        self.client.force_login(self.student)
        response = self.client.post(quiz.get_absolute_url(), {'reponse_0': 'B'})
        self.assertNotIn('score', response.context)
        self.assertFalse(response.context.get('corrige', False))
        self.assertIn('erreur_reponse', response.context)
        response = self.client.post(quiz.get_absolute_url(), {'reponse_0': 'A', 'reponse_1': 'D'})
        self.assertEqual(response.context['score'], 2)
        self.client.force_login(self.user)
        response = self.client.get(reverse('pedagogie:exercice_update', args=[quiz.pk]))
        self.assertContains(response, 'Second?')
        self.data.update({'questions-INITIAL_FORMS': '1', 'questions-0-DELETE': 'on'})
        self.client.post(reverse('pedagogie:exercice_update', args=[quiz.pk]), self.data)
        quiz.refresh_from_db()
        self.assertEqual(quiz.questions_supplementaires, [])

    def test_invalid_extra_question_does_not_save(self):
        self.data['questions-0-bonne_reponse'] = 'invalid'
        response = self.client.post(reverse('pedagogie:exercice_create'), self.data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Exercice.objects.exists())
        self.assertTrue(response.context['questions_formset'].errors)

    def test_existing_single_question(self):
        quiz = Exercice.objects.create(cours=self.course, titre='Old', question='First?', propositions='A\nB', bonne_reponse='A')
        self.client.force_login(self.student)
        response = self.client.post(quiz.get_absolute_url(), {'reponse_0': 'A'})
        self.assertEqual(response.context['score'], 1)

    def test_multiple_true_false_and_open_questions(self):
        for kind, first, second in [('vf', 'Vrai', 'Faux'), ('ouverte', 'Hello', 'Goodbye')]:
            with self.subTest(kind=kind):
                self.data.update(type_exercice=kind, bonne_reponse=first)
                self.data['questions-0-bonne_reponse'] = second
                self.data.pop('propositions', None)
                self.data.pop('questions-0-propositions', None)
                response = self.client.post(reverse('pedagogie:exercice_create'), self.data)
                self.assertEqual(response.status_code, 302)
                quiz = Exercice.objects.latest('pk')
                self.assertEqual(len(quiz.questions_supplementaires), 1)
                self.client.force_login(self.student)
                response = self.client.post(quiz.get_absolute_url(), {'reponse_0': first, 'reponse_1': second})
                self.assertEqual(response.context['score'], None if kind == 'ouverte' else 2)
                self.client.force_login(self.user)
                response = self.client.get(reverse('pedagogie:exercice_update', args=[quiz.pk]))
                self.assertContains(response, 'Second?')
                self.data['questions-0-bonne_reponse'] = ''
                response = self.client.post(reverse('pedagogie:exercice_create'), self.data)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['questions_formset'].errors)


class CourseChapterTests(TestCase):
    def setUp(self):
        self.teacher = get_user_model().objects.create_superuser(username='teacher', password='test')
        self.client.force_login(self.teacher)
        self.data = {'titre': 'English', 'description': 'A course', 'matiere': 'English', 'niveau': 'debutant',
                     'contenu': '<p><strong>Welcome</strong></p>',
                     'chapitres-TOTAL_FORMS': '2', 'chapitres-INITIAL_FORMS': '0',
                     'chapitres-0-titre': 'Greetings', 'chapitres-0-contenu': '<h2>Hello</h2><font color="#ff0000">World</font>',
                     'chapitres-1-titre': 'Grammar', 'chapitres-1-contenu': '<ul><li>Verbs</li></ul>'}

    def test_save_read_as_learner_and_edit(self):
        response = self.client.post(reverse('pedagogie:cours_create'), self.data)
        self.assertEqual(response.status_code, 302)
        course = Cours.objects.get()
        self.assertEqual(course.enseignant, self.teacher)
        self.assertEqual(len(course.chapitres), 2)
        self.assertIn('<strong>', course.contenu)
        self.assertIn('color:#ff0000', course.chapitres[0]['contenu'])
        response = self.client.get(reverse('pedagogie:cours_update', args=[course.pk]))
        self.assertContains(response, 'Greetings')
        self.data.update({'chapitres-INITIAL_FORMS': '2', 'chapitres-0-DELETE': 'on'})
        self.client.post(reverse('pedagogie:cours_update', args=[course.pk]), self.data)
        course.refresh_from_db()
        self.assertEqual([c['titre'] for c in course.chapitres], ['Grammar'])
        from accounts.roles import ensure_system_roles, LEARNER
        from django.contrib.auth.models import Group
        ensure_system_roles()
        learner = get_user_model().objects.create_user(username='learner')
        learner.groups.add(Group.objects.get(name=LEARNER))
        self.client.force_login(learner)
        response = self.client.get(reverse('pedagogie:cours_list'))
        self.assertContains(response, 'English')
        response = self.client.get(course.get_absolute_url())
        self.assertContains(response, 'Sommaire')
        response = self.client.get(course.get_absolute_url(), {'chapitre': course.chapitres[0]['id']})
        self.assertContains(response, 'Verbs')
        self.assertContains(response, 'reading-bullet')
        self.assertNotContains(response, 'Retirer ce chapitre')
        self.assertEqual(self.client.post(reverse('pedagogie:cours_update', args=[course.pk]), self.data).status_code, 403)

    def test_invalid_chapter_and_safe_formatting(self):
        self.data['chapitres-1-titre'] = ''
        response = self.client.post(reverse('pedagogie:cours_create'), self.data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Cours.objects.exists())
        from pedagogie.richtext import clean_content
        result = clean_content('<img src=x onerror=alert(1)><p onclick="bad()" style="color:red;background-image:url(x)"><b>Safe</b></p><script>alert(1)</script>')
        self.assertIn('<b>Safe</b>', result)
        self.assertIn('color:red', result)
        for unsafe in ['<script', 'onerror', 'onclick', 'background-image']:
            self.assertNotIn(unsafe, result)


class CourseImportTests(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        settings = override_settings(MEDIA_ROOT=self.directory.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.teacher = get_user_model().objects.create_superuser(username='teacher', password='test')
        self.client.force_login(self.teacher)

    def word(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Introduction')
        document.add_heading('Chapter A', 1)
        document.add_paragraph().add_run('Important').bold = True
        document.add_heading('Chapter B', 1)
        document.add_paragraph('Second chapter')
        stream = BytesIO()
        document.save(stream)
        return SimpleUploadedFile('Word course.docx', stream.getvalue())

    def slides(self):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = 'Slide A'
        slide.placeholders[1].text = 'Slide text'
        stream = BytesIO()
        presentation.save(stream)
        return SimpleUploadedFile('Slides.pptx', stream.getvalue())

    def test_bulk_import_download_and_exports(self):
        response = self.client.post(reverse('pedagogie:cours_import'), {
            'matiere': 'English', 'niveau': 'debutant', 'fichiers': [self.word(), self.slides()]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Cours.objects.count(), 2)
        self.assertEqual(SupportCours.objects.count(), 2)
        course = Cours.objects.get(titre='Word course')
        self.assertEqual(len(course.chapitres), 2)
        self.assertIn('<strong>Important</strong>', course.chapitres[0]['contenu'])
        slides = Cours.objects.get(titre='Slides')
        self.assertEqual(slides.chapitres[0]['titre'], 'Slide A')
        self.assertIn('Slide text', slides.chapitres[0]['contenu'])
        support = course.supports.get()
        response = self.client.get(reverse('pedagogie:support_download', args=[support.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertTrue(b''.join(response.streaming_content).startswith(b'PK'))
        from docx import Document
        from pptx import Presentation
        for kind, reader in [('docx', Document), ('pptx', Presentation)]:
            response = self.client.get(reverse('pedagogie:cours_export', args=[course.pk, kind]))
            self.assertEqual(response.status_code, 200)
            document = reader(BytesIO(b''.join(response.streaming_content)))
            self.assertIsNotNone(document)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('pedagogie:support_download', args=[support.pk])).status_code, 302)

    def test_bad_batch_creates_nothing(self):
        response = self.client.post(reverse('pedagogie:cours_import'), {
            'matiere': 'English', 'niveau': 'debutant',
            'fichiers': [self.word(), SimpleUploadedFile('broken.docx', b'not a document')]})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['form'].errors)
        self.assertFalse(Cours.objects.exists())
        self.assertFalse(SupportCours.objects.exists())

    def test_attach_multiple_files_to_existing_course(self):
        course = Cours.objects.create(titre='Existing', enseignant=self.teacher)
        response = self.client.post(reverse('pedagogie:cours_update', args=[course.pk]), {
            'titre': 'Existing', 'description': 'Description', 'matiere': 'English', 'niveau': 'debutant',
            'contenu': 'Welcome', 'chapitres-TOTAL_FORMS': '0', 'chapitres-INITIAL_FORMS': '0',
            'fichiers': [self.word(), self.slides()]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(course.supports.count(), 2)
        self.assertEqual(Cours.objects.count(), 1)

    def test_chapter_scope_permissions_and_stable_identity(self):
        from accounts.roles import ensure_system_roles, TEACHER
        from django.contrib.auth.models import Group
        ensure_system_roles()
        teacher = get_user_model().objects.create_user(username='owner')
        teacher.groups.add(Group.objects.get(name=TEACHER))
        first, second = str(uuid4()), str(uuid4())
        course = Cours.objects.create(titre='Course', enseignant=teacher,
            chapitres=[{'id': first, 'titre': 'First', 'contenu': 'First body'}, {'id': second, 'titre': 'Second', 'contenu': 'Second body'}])
        self.client.force_login(teacher)
        data = {'titre': 'Quiz', 'cours': course.pk, 'chapitre_id': second,
            'type_exercice': 'vf', 'difficulte': 'facile', 'question': 'True?', 'bonne_reponse': 'Vrai',
            'questions-TOTAL_FORMS': '0', 'questions-INITIAL_FORMS': '0'}
        response = self.client.post(reverse('pedagogie:exercice_create'), data)
        self.assertEqual(response.status_code, 302)
        exercise = Exercice.objects.get()
        self.assertEqual(exercise.chapitre_id, second)
        response = self.client.get(course.get_absolute_url(), {'chapitre': first})
        self.assertNotIn(exercise, response.context['exercices_affiches'])
        response = self.client.get(course.get_absolute_url(), {'chapitre': second})
        self.assertIn(exercise, response.context['exercices_affiches'])
        response = self.client.get(reverse('pedagogie:exercice_create'), {'cours': course.pk, 'chapitre': second})
        self.assertEqual(response.context['form'].initial['chapitre_id'], second)
        data['chapitre_id'] = str(uuid4())
        self.assertEqual(self.client.post(reverse('pedagogie:exercice_create'), data).status_code, 200)
        foreign = Cours.objects.create(titre='Other', enseignant=self.teacher)
        self.assertEqual(self.client.get(reverse('pedagogie:chapitres_choices', args=[foreign.pk])).status_code, 404)
        data['cours'] = foreign.pk
        self.assertEqual(self.client.post(reverse('pedagogie:exercice_create'), data).status_code, 200)
        # Delete the first chapter and rename the second: identity must survive.
        payload = {'titre': 'Course', 'description': 'Description', 'matiere': 'English', 'niveau': 'debutant', 'contenu': 'Intro',
            'chapitres-TOTAL_FORMS': '2', 'chapitres-INITIAL_FORMS': '2',
            'chapitres-0-id': first, 'chapitres-0-DELETE': 'on',
            'chapitres-1-id': second, 'chapitres-1-titre': 'Renamed', 'chapitres-1-contenu': 'Body'}
        self.assertEqual(self.client.post(reverse('pedagogie:cours_update', args=[course.pk]), payload).status_code, 302)
        exercise.refresh_from_db()
        self.assertEqual(exercise.chapitre_id, second)
        self.assertEqual(exercise.chapitre_titre, 'Renamed')
        payload.update({'chapitres-1-DELETE': 'on'})
        self.client.post(reverse('pedagogie:cours_update', args=[course.pk]), payload)
        exercise.refresh_from_db()
        self.assertEqual(exercise.chapitre_id, '')
