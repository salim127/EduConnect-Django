"""Structured, paginated PDF and slide exports (no raw HTML or carriage returns)."""
from functools import lru_cache
from io import BytesIO
import ipaddress
import socket
from urllib.parse import urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener

from PIL import Image
from .presentation import reading_pages


def public_url(url):
    parts = urlsplit(url)
    if parts.scheme not in ('https', 'http') or not parts.hostname or parts.username or parts.password:
        raise ValueError('Only public HTTP images can be exported')
    if parts.port not in (None, 80, 443):
        raise ValueError('Unsupported image port')
    addresses = socket.getaddrinfo(parts.hostname, parts.port or (443 if parts.scheme == 'https' else 80), type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(address[4][0]).is_global for address in addresses):
        raise ValueError('Private image addresses are not supported')


class PublicRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


@lru_cache(maxsize=16)
def image_data(url):
    """Small raster images only; never fetch private network or local file URLs."""
    try:
        public_url(url)
        request = Request(url, headers={'User-Agent': 'EduConnect/1.0'})
        with build_opener(PublicRedirect()).open(request, timeout=6) as response:
            raw = response.read(8 * 1024 * 1024 + 1)
        if len(raw) > 8 * 1024 * 1024:
            return None
        with Image.open(BytesIO(raw)) as image:
            if image.width * image.height > 20_000_000:
                return None
            image.thumbnail((1600, 1000))
            image = image.convert('RGB')
            buffer = BytesIO()
            image.save(buffer, 'PNG')
            return buffer.getvalue()
    except (OSError, ValueError, Image.DecompressionBombError):
        return None


def course_pages(course):
    sections = [('Introduction', course.contenu)] + [(c['titre'], c['contenu']) for c in course.chapitres]
    return [(title, number, blocks) for title, html in sections for number, blocks in enumerate(reading_pages(html, capacity=15), 1)]


def fit_image(raw, width, height):
    with Image.open(BytesIO(raw)) as image:
        ratio = min(width / image.width, height / image.height)
        return image.width * ratio, image.height * ratio


def pptx_export(course, output):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    purple, ink = RGBColor.from_string('493B96'), RGBColor.from_string('25204A')

    def textbox(slide, x, y, w, h, text, size, color=ink, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(text.splitlines() or ['']):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = 'Aptos'
            paragraph.font.size, paragraph.font.color.rgb, paragraph.font.bold = Pt(size), color, bold
            paragraph.space_after = Pt(5)
        return box

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = purple
    white = RGBColor(255, 255, 255)
    textbox(cover, .8, 1.5, 11.7, 2.1, course.titre, 36, white, True)
    textbox(cover, .8, 4, 11.7, 1.5, course.description, 21, white)
    pages = course_pages(course)
    for index, (title, number, blocks) in enumerate(pages, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox(slide, .65, .3, 12, .7, title + (f' — {number}' if number > 1 else ''), 26, purple, True)
        textbox(slide, .65, 7.12, 11, .25, course.titre[:110], 10)
        textbox(slide, 12, 7.12, .7, .25, f'{index}/{len(pages)}', 10)
        y = 1.25
        unit = 5.5 / max(15, sum(b['units'] for b in blocks))
        for block in blocks:
            height = block['units'] * unit
            if block['kind'] == 'image':
                raw = image_data(block['src'])
                if raw:
                    width, actual_height = fit_image(raw, 11.8, height - .12)
                    slide.shapes.add_picture(BytesIO(raw), Inches(.75 + (11.8 - width) / 2), Inches(y), width=Inches(width), height=Inches(actual_height))
                else:
                    textbox(slide, .75, y, 11.8, height, 'Image non disponible hors connexion : ' + block['alt'][:180], 16)
            elif block['kind'] == 'table':
                rows = block['rows']
                columns = max(len(row) for row in rows)
                table = slide.shapes.add_table(len(rows), columns, Inches(.75), Inches(y), Inches(11.8), Inches(height - .12)).table
                for row_index, row in enumerate(rows):
                    for column_index, cell in enumerate(row):
                        target = table.cell(row_index, column_index)
                        target.text = cell
                        target.text_frame.word_wrap = True
                        for paragraph in target.text_frame.paragraphs:
                            paragraph.font.size = Pt(16)
                            paragraph.font.color.rgb = ink
            else:
                text = ('• ' if block['kind'] == 'bullet' else '') + block['text']
                textbox(slide, .75, y, 11.8, height, text, 23 if block['kind'] == 'heading' else 18,
                        purple if block['kind'] == 'heading' else ink, block['kind'] == 'heading')
            y += height
    presentation.save(output)


def pdf_export(course, output):
    from .document_export import browser_pdf
    from .pdf_document import pdf_document
    if not browser_pdf(course, output):
        pdf_document(course, output)
