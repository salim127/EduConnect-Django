"""Local Word/PowerPoint import and export; no Microsoft account required."""
from html import escape
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from django.core.exceptions import ValidationError

MAX_FILE_SIZE = 15 * 1024 * 1024


def validate_support(upload):
    if upload.size > MAX_FILE_SIZE:
        raise ValidationError('Chaque fichier doit faire au maximum 15 Mo.')
    if Path(upload.name).suffix.lower() != '.pdf':
        read_office(upload)
        return
    from pypdf import PdfReader
    try:
        upload.seek(0)
        if upload.read(5) != b'%PDF-':
            raise ValueError('Invalid PDF signature')
        upload.seek(0)
        reader = PdfReader(upload)
        if reader.is_encrypted or not len(reader.pages):
            raise ValueError('Encrypted or empty PDF')
    except Exception as exc:
        raise ValidationError('PDF illisible ou protégé. Choisissez un PDF valide, sans mot de passe.') from exc
    finally:
        upload.seek(0)


def chapter(title, content):
    return {'id': str(uuid4()), 'titre': title[:200], 'contenu': content}


def read_office(upload):
    suffix = Path(upload.name).suffix.lower()
    if suffix not in ('.docx', '.pptx'):
        raise ValidationError('Formats acceptés : Word (.docx) et PowerPoint (.pptx).')
    if upload.size > MAX_FILE_SIZE:
        raise ValidationError('Chaque fichier doit faire au maximum 15 Mo.')
    upload.seek(0)
    data = upload.read()
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 3000 or sum(i.file_size for i in infos) > 60 * 1024 * 1024:
                raise ValidationError('Le document est trop volumineux une fois décompressé.')
            expected = 'word/document.xml' if suffix == '.docx' else 'ppt/presentation.xml'
            if expected not in archive.namelist():
                raise ValidationError('Le contenu du fichier ne correspond pas à son extension.')
        if suffix == '.docx':
            result = read_word(data)
        else:
            result = read_powerpoint(data)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError('Document illisible ou protégé. Enregistrez-le à nouveau en .docx ou .pptx.') from exc
    finally:
        upload.seek(0)
    return result


def word_paragraph(paragraph):
    parts = []
    for run in paragraph.runs:
        text = escape(run.text).replace('\n', '<br>')
        if run.bold:
            text = f'<strong>{text}</strong>'
        if run.italic:
            text = f'<em>{text}</em>'
        if run.underline:
            text = f'<u>{text}</u>'
        if run.font.color.rgb is not None:
            text = f'<span style="color:#{run.font.color.rgb}">{text}</span>'
        parts.append(text)
    return '<p>' + ''.join(parts) + '</p>'


def read_word(data):
    from docx import Document
    document = Document(BytesIO(data))
    intro, chapters = [], []
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        style = paragraph.style.name if paragraph.style else ''
        if style in ('Heading 1', 'Heading 2', 'Titre 1', 'Titre 2'):
            chapters.append(chapter(paragraph.text, ''))
        elif chapters:
            chapters[-1]['contenu'] += word_paragraph(paragraph)
        else:
            intro.append(word_paragraph(paragraph))
    if not intro and not chapters:
        raise ValidationError('Ce document ne contient pas de texte importable.')
    if len(chapters) > 100:
        raise ValidationError('Un cours peut contenir au maximum 100 chapitres.')
    return ''.join(intro) or '<p>Consultez les chapitres ci-dessous.</p>', chapters


def read_powerpoint(data):
    from pptx import Presentation
    presentation = Presentation(BytesIO(data))
    if not presentation.slides or len(presentation.slides) > 100:
        raise ValidationError('Le diaporama doit contenir entre 1 et 100 diapositives.')
    chapters = []
    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text if title_shape is not None and title_shape.text.strip() else f'Diapositive {number}'
        content = []
        for shape in slide.shapes:
            if shape == title_shape or not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text.strip():
                    content.append('<p>' + escape(paragraph.text) + '</p>')
        chapters.append(chapter(title, ''.join(content) or '<p>Voir le fichier original pour les éléments visuels.</p>'))
    return '<p>Support de présentation : un chapitre par diapositive.</p>', chapters


def plain_text(html):
    from .presentation import content_blocks
    return '\n'.join(block['text'] for block in content_blocks(html) if block['kind'] != 'image')


def export_course(course, kind):
    output = BytesIO()
    if kind == 'pdf':
        from .export_layout import pdf_export
        pdf_export(course, output)
    elif kind == 'docx':
        from .document_export import word_export
        word_export(course, output)
    else:
        from .export_layout import pptx_export
        pptx_export(course, output)
    output.seek(0)
    return output
