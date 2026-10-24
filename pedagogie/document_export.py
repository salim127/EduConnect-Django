"""Document exports preserve editor formatting and use natural document flow."""
from copy import copy
from html import escape
from io import BytesIO
from pathlib import Path
import base64
import os
import re
import shutil
import subprocess
import tempfile

from bs4 import BeautifulSoup, NavigableString
from .presentation import normalized_text
from .richtext import clean_content
from . import export_layout as assets


def css(node):
    return dict((key.strip().lower(), value.strip()) for part in node.get('style', '').split(';') if ':' in part for key, value in [part.split(':', 1)])


def points(value, base=12):
    match = re.fullmatch(r'([0-9.]+)(px|pt|em|rem|%)?', str(value or '').strip())
    if not match:
        return None
    number, unit = float(match[1]), match[2]
    return number * ({'px': .75, 'pt': 1, 'em': base, 'rem': 12, '%': base / 100}.get(unit, 1))


def rich_blocks(html):
    """Keep original paragraph boundaries and inline markup, including nested pasted HTML."""
    soup = BeautifulSoup(clean_content(normalized_text(html)), 'html.parser')
    result = []
    boundaries = {'div', 'p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'ul', 'ol', 'li', 'table', 'img', 'pre', 'blockquote'}

    def walk(container, prefix='', inherited=None):
        style = dict(inherited or {})
        style.update(css(container))
        pending = []

        def flush():
            nonlocal prefix
            if not pending:
                return
            paragraph = soup.new_tag('p')
            for item in pending:
                paragraph.append(copy(item))
            pending.clear()
            if paragraph.get_text().strip() or paragraph.find('br'):
                result.append(('paragraph', paragraph, prefix, style))
                prefix = ''

        for child in container.children:
            if isinstance(child, NavigableString):
                pending.append(child)
            elif child.name not in boundaries and not child.find(list(boundaries)):
                pending.append(child)
            else:
                flush()
                kind = child.name
                child_style = dict(style)
                child_style.update(css(child))
                if kind in ('table', 'img', 'pre'):
                    result.append((kind, child, '', child_style))
                elif re.fullmatch(r'h[1-6]', kind) or kind == 'p' and not child.find(list(boundaries)):
                    result.append((kind, child, prefix if kind == 'p' else '', child_style))
                    prefix = ''
                elif kind in ('ul', 'ol'):
                    for index, item in enumerate(child.find_all('li', recursive=False), 1):
                        walk(item, f'{index}.' if kind == 'ol' else '•', child_style)
                else:
                    walk(child, prefix, child_style)
                    prefix = ''
        flush()

    walk(soup)
    return result


def word_export(course, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    style = document.styles['Normal']
    style.font.name, style.font.size = 'Arial', Pt(12)
    style.paragraph_format.line_spacing = 1.7
    style.paragraph_format.space_after = Pt(12)
    document.add_heading(course.titre, 0)
    if course.description.strip():
        document.add_paragraph(course.description)

    def write_inline(paragraph, node, properties=None, pre=False):
        properties = dict(properties or {})
        if isinstance(node, NavigableString):
            text = normalized_text(str(node))
            if not pre:
                text = re.sub(r'\s+', ' ', text)
                if paragraph.runs and paragraph.runs[-1].text.endswith(' '):
                    text = text.lstrip(' ')
            if not text:
                return
            run = paragraph.add_run(text)
            for name in ('bold', 'italic', 'underline'):
                if name in properties:
                    setattr(run, name, properties[name])
            if properties.get('strike'):
                run.font.strike = True
            if properties.get('font-weight') in ('bold', '600', '700', '800', '900'):
                run.bold = True
            if properties.get('font-style') == 'italic':
                run.italic = True
            if 'underline' in properties.get('text-decoration', ''):
                run.underline = True
            if properties.get('font-family'):
                run.font.name = properties['font-family'].split(',')[0].strip(' \"\'')
            size = points(properties.get('font-size'))
            if size:
                run.font.size = Pt(min(96, size))
            color = properties.get('color', '').lstrip('#')
            names = {'red': 'FF0000', 'blue': '0000FF', 'green': '008000', 'black': '000000', 'white': 'FFFFFF'}
            color = names.get(color.lower(), color)
            if re.fullmatch(r'[0-9a-fA-F]{3}', color):
                color = ''.join(char * 2 for char in color)
            if re.fullmatch(r'[0-9a-fA-F]{6}', color):
                run.font.color.rgb = RGBColor.from_string(color.upper())
            return
        properties.update(css(node))
        if node.name in ('b', 'strong'):
            properties['bold'] = True
        if node.name in ('i', 'em'):
            properties['italic'] = True
        if node.name == 'u':
            properties['underline'] = True
        if node.name in ('s', 'del'):
            properties['strike'] = True
        if node.name == 'font' and node.get('color'):
            properties['color'] = node['color']
        if node.name == 'br':
            paragraph.add_run().add_break()
            return
        for child in node.children:
            write_inline(paragraph, child, properties, pre)

    for chapter_title, html in [('Introduction', course.contenu)] + [(c['titre'], c['contenu']) for c in course.chapitres]:
        document.add_heading(chapter_title, 1)
        for kind, node, prefix, properties in rich_blocks(html):
            if kind == 'img':
                raw = assets.image_data(node.get('src', ''))
                if raw:
                    width, height = assets.fit_image(raw, 6.77, 8.8)
                    picture = document.add_paragraph()
                    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    picture.add_run().add_picture(BytesIO(raw), width=Inches(width), height=Inches(height))
                else:
                    document.add_paragraph(node.get('alt') or 'Image distante indisponible')
            elif kind == 'table':
                rows = node.find_all('tr')
                count = max((len(row.find_all(['td', 'th'], recursive=False)) for row in rows), default=0)
                if not count:
                    continue
                table = document.add_table(rows=len(rows), cols=count)
                table.style = 'Table Grid'
                for row_index, row in enumerate(rows):
                    for column, cell in enumerate(row.find_all(['td', 'th'], recursive=False)):
                        paragraph = table.cell(row_index, column).paragraphs[0]
                        for child in cell.children:
                            write_inline(paragraph, child, {'bold': cell.name == 'th'})
            else:
                paragraph = document.add_heading('', min(6, int(kind[1]))) if re.fullmatch(r'h[1-6]', kind) else document.add_paragraph()
                paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                                       'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}.get(properties.get('text-align'), WD_ALIGN_PARAGRAPH.LEFT)
                line_height = properties.get('line-height', '')
                if re.fullmatch(r'[0-9.]+', line_height):
                    paragraph.paragraph_format.line_spacing = max(.8, min(3, float(line_height)))
                elif points(line_height):
                    paragraph.paragraph_format.line_spacing = Pt(points(line_height))
                if prefix:
                    paragraph.add_run(prefix + ' ')
                if kind == 'pre':
                    properties['font-family'] = 'Consolas'
                for child in node.children:
                    write_inline(paragraph, child, properties, pre=kind == 'pre')
    document.save(output)


def browser_pdf(course, output):
    """Print sanitized, self-contained editor HTML with the browser's natural layout."""
    candidates = [os.environ.get('COURSE_PDF_BROWSER', ''), shutil.which('google-chrome'), shutil.which('chromium'), shutil.which('chromium-browser'),
                  str(Path(os.environ.get('PROGRAMFILES', 'C:/Program Files')) / 'Google/Chrome/Application/chrome.exe'),
                  str(Path(os.environ.get('PROGRAMFILES(X86)', 'C:/Program Files (x86)')) / 'Microsoft/Edge/Application/msedge.exe')]
    browser = next((candidate for candidate in candidates if candidate and Path(candidate).is_file()), None)
    if not browser:
        return False
    sections = []
    for title, html in [('Introduction', course.contenu)] + [(c['titre'], c['contenu']) for c in course.chapitres]:
        soup = BeautifulSoup(clean_content(normalized_text(html)), 'html.parser')
        for image in soup.find_all('img'):
            raw = assets.image_data(image.get('src', ''))
            if raw:
                image['src'] = 'data:image/png;base64,' + base64.b64encode(raw).decode()
            else:
                image.replace_with(soup.new_string(image.get('alt') or 'Image distante indisponible'))
        sections.append('<section><h1 class="chapter-title">' + escape(title) + '</h1>' + str(soup) + '</section>')
    html = '''<!doctype html><html lang="fr"><meta charset="utf-8"><style>
      @page{size:A4;margin:18mm}body{font-family:Arial,sans-serif;font-size:12pt;line-height:1.7;color:#25204a}
      h1{font-size:22pt}h2{font-size:18pt}h3{font-size:15pt}h1,h2,h3,h4{break-after:avoid;margin:16pt 0 8pt}
      p{margin:12pt 0}ul,ol{margin:6pt 0;padding-left:22pt}li{margin-bottom:4pt}
      .chapter-title{font-size:20pt;color:#493b96}
      img{max-width:100%!important;height:auto!important;max-height:235mm;object-fit:contain;break-inside:avoid}
      table{border-collapse:collapse;max-width:100%;width:100%;table-layout:fixed}td,th{border:1px solid #dcd8ed;padding:6pt;overflow-wrap:anywhere}
      tr{break-inside:avoid}thead{display:table-header-group}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:10pt}
      a{color:#493b96}body>header{margin-bottom:18pt;border-bottom:1px solid #dcd8ed;padding-bottom:10pt}
      </style><body><header><h1>''' + escape(course.titre) + '</h1><p>' + escape(course.description) + '</p></header>' + ''.join(sections) + '</body></html>'
    try:
        with tempfile.TemporaryDirectory(prefix='edu-pdf-', ignore_cleanup_errors=True) as directory:
            folder = Path(directory)
            source, target = folder / 'course.html', folder / 'course.pdf'
            source.write_text(html, encoding='utf-8')
            command = [browser, '--headless=new', '--disable-gpu', '--disable-extensions', '--disable-background-networking',
                       '--no-first-run', '--no-default-browser-check', '--no-pdf-header-footer', '--host-resolver-rules=MAP * ~NOTFOUND',
                       f'--user-data-dir={folder / "profile"}', f'--print-to-pdf={target}', source.as_uri()]
            result = subprocess.run(command, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=35,
                                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            if result.returncode or not target.exists():
                return False
            raw = target.read_bytes()
            if not raw.startswith(b'%PDF'):
                return False
            output.write(raw)
            return True
    except (OSError, subprocess.TimeoutExpired):
        return False
