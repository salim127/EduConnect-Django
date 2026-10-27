from io import BytesIO
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.test import TestCase, SimpleTestCase
from django.urls import reverse
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from .models import Cours
from .office import export_course, plain_text
from .presentation import content_blocks, reading_pages


class PresentationContentTests(SimpleTestCase):
    def test_nested_pasted_html_is_reflowed_without_duplicate_text(self):
        html = '<ul><li>First line_x000D_<h1>Heading</h1><p>Paragraph</p><div><img src="https://example.com/a.png" width="3000"></div><div>·</div><h2>Next</h2><table><tr><td>Cell</td><td>Value</td></tr></table></li></ul>'
        blocks = content_blocks(html)
        self.assertEqual([b['kind'] for b in blocks], ['bullet', 'heading', 'paragraph', 'image', 'heading', 'table'])
        self.assertEqual([b['text'] for b in blocks].count('Heading'), 1)
        self.assertNotIn('_x000D_', plain_text(html))
        self.assertIn('Cell | Value', plain_text(html))
        self.assertGreater(len(reading_pages(html)), 1)

    def test_long_content_keeps_all_words_on_bounded_pages(self):
        text = ' '.join(f'word{index}' for index in range(800))
        pages = reading_pages('<p>' + text + '</p>')
        self.assertGreater(len(pages), 2)
        result = ' '.join(block['text'] for page in pages for block in page)
        self.assertEqual(result.split(), text.split())
        self.assertTrue(all(sum(b['units'] for b in page) <= 22 for page in pages))

    def test_reader_keeps_image_with_short_opening_text(self):
        html = '<p>Opening line</p><h1>Introduction au développement Web moderne</h1><p>Cours complet HTML, CSS et JavaScript.</p><img src="https://example.com/course.png"><p>Caption</p><h2>Next section</h2><p>' + 'Following content. ' * 30 + '</p>'
        pages = reading_pages(html)
        self.assertIn('image', [block['kind'] for block in pages[0]])
        self.assertIn('Caption', [block['text'] for block in pages[0]])
        self.assertTrue(all(sum(block['units'] for block in page) <= 15 for page in reading_pages(html, capacity=15)))


class PresentationExportTests(TestCase):
    def setUp(self):
        teacher = get_user_model().objects.create_superuser(username='teacher', password='test')
        self.client.force_login(teacher)
        self.course = Cours.objects.create(titre='Web Course', description='Learn the web', enseignant=teacher,
            contenu='<h1>Introduction</h1><p>Hello_x000D_ world</p><img src="https://example.com/image.png" alt="Illustration"><h2>Technologies</h2><table><tr><td>HTML</td><td>Structure</td></tr><tr><td>CSS</td><td>Style</td></tr></table>' + '<p>Learning content for the next page. </p>' * 30)
        buffer = BytesIO()
        Image.new('RGB', (1800, 800), 'blue').save(buffer, 'PNG')
        self.image = buffer.getvalue()

    def test_pdf_and_pptx_preserve_images_tables_and_bounds(self):
        with patch('pedagogie.export_layout.image_data', return_value=self.image), patch('pedagogie.document_export.browser_pdf', return_value=False):
            pdf = export_course(self.course, 'pdf')
            pptx = Presentation(export_course(self.course, 'pptx'))
        reader = PdfReader(pdf)
        pdf_text = ''.join(page.extract_text() for page in reader.pages)
        self.assertNotIn('_x000D_', pdf_text)
        self.assertIn('HTML', pdf_text)
        self.assertTrue(any(page.images for page in reader.pages))
        self.assertGreaterEqual(len(reader.pages), 1)
        text = '\n'.join(shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertNotIn('_x000D_', text)
        self.assertTrue(any(shape.has_table for slide in pptx.slides for shape in slide.shapes))
        pictures = [shape for slide in pptx.slides for shape in slide.shapes if shape.shape_type == 13]
        self.assertTrue(pictures)
        for slide in pptx.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, pptx.slide_width)
                self.assertLessEqual(shape.top + shape.height, pptx.slide_height)

    def test_reader_pagination_and_pdf_download(self):
        first = self.client.get(self.course.get_absolute_url())
        self.assertGreater(first.context['pages_lecture'].num_pages, 1)
        self.assertContains(first, 'Page suivante')
        self.assertContains(first, 'exporter/pdf/')
        second = self.client.get(self.course.get_absolute_url(), {'page': 2})
        self.assertEqual(second.context['page_lecture'].number, 2)
        self.assertNotEqual(first.context['page_lecture'].object_list, second.context['page_lecture'].object_list)
        with patch('pedagogie.export_layout.image_data', return_value=self.image), patch('pedagogie.document_export.browser_pdf', return_value=False):
            response = self.client.get(reverse('pedagogie:cours_export', args=[self.course.pk, 'pdf']))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response['Content-Type'], 'application/pdf')
            self.assertTrue(b''.join(response.streaming_content).startswith(b'%PDF'))

    def test_export_never_fetches_private_or_file_images(self):
        from .export_layout import image_data
        for url in ('file:///etc/passwd', 'http://127.0.0.1/secret', 'http://169.254.169.254/metadata'):
            with patch('pedagogie.export_layout.build_opener') as opener:
                self.assertIsNone(image_data(url))
                opener.assert_not_called()
